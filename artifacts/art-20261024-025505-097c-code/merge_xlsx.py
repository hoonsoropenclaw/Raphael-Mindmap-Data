#!/usr/bin/env python3
"""
合併資料夾內多個 Excel 並產生摘要（v2 改良版）
==================================================

用途（admin 軌：辦公自動化，Windows 使用者環境）：
  給定一個資料夾，把裡面所有 .xlsx 合併成單一檔案，並產生摘要工作表：
    - 合併資料（跨檔聯集欄位，缺漏欄位留白）
    - 檔案摘要（每檔每分頁的列數、欄位、缺漏／獨有欄位、錯誤訊息）
    - 數值合計（每欄合計 + 非空數值格數）
    - 退款/負值警示（欄名關鍵字 或 數值為負）
    - 本次執行（時間、輸入、輸出、統計）

與 v1（2026-10-07 / 2026-10-16）的差異：
  1. 輸出採「暫存檔 + fsync + os.replace」原子寫入（v1 直接 wb.save 覆寫，中途失敗會留半套檔）
  2. 既有輸出檔改存「時間戳備份」，不再用固定 .bak（v1 的 .bak 每次被覆蓋＝沒有版本）
  3. 新增 --summary-csv：輸出摘要 CSV，預設 utf-8-sig（Excel 直接開不亂碼），可切 cp950
  4. 新增 Windows 路徑預檢：非法字元 / 保留裝置名 / 超過 MAX_PATH 260，先擋再動磁碟
  5. 跨檔欄位聯集 + 缺漏/獨有欄位報告（v1 是 pd.concat，欄位不齊時只靜默補 NaN）
  6. 僅依賴 openpyxl（v1 還要 pandas），Windows 端只需 pip install openpyxl
  7. 錯誤處理：壞檔略過不中斷、錯誤訊息可行動（說明要改什麼）

執行環境：
  - 跨平台（Windows / macOS / Linux），Python 3.10+
  - 安裝：pip install openpyxl

執行：
  python merge_xlsx.py <input_folder>                 # 預設 dry-run，只預覽不寫檔
  python merge_xlsx.py <input_folder> --apply         # 真的合併並寫出
  python merge_xlsx.py <input_folder> --apply \
      -o out/merged.xlsx --summary-csv out/summary.csv --csv-encoding cp950

設計原則（sops/code/SOP.md）：
  - 破壞性動作（寫檔/覆寫）預設 dry-run，需 --apply 才真的執行
  - 任何外部 I/O 失敗都要有可行動的錯誤訊息
  - 可重跑不損壞資料（idempotent：重跑覆寫同一份輸出，資料內容一致）
"""
from __future__ import annotations

import argparse
import csv
import os
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path, PureWindowsPath

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - 環境缺套件時給可行動訊息
    sys.stderr.write(
        "缺少必要套件 openpyxl。請執行：pip install openpyxl\n"
        "（或用 python -m pip install --user openpyxl）\n"
    )
    raise SystemExit(3)


# ---------- 設定 ----------
REFUND_KEYWORDS = ("退款", "退費", "折讓", "溢付", "沖銷", "refund", "credit")
INPUT_SUFFIXES = {".xlsx"}
# 其他 Excel 家族格式（會被回報略過原因）。.csv 不列入——它不是 Excel 檔，
# 在「略過的 Excel 檔案」清單裡出現只會造成噪音。
OTHER_EXCEL_SUFFIXES = {".xls", ".xlsm", ".xlsb"}
SKIP_PREFIXES = ("~$", ".")  # Excel 鎖檔 / 隱藏檔
CSV_ENCODINGS = ("utf-8-sig", "cp950", "utf-8")
MAX_WINDOWS_PATH = 260
ILLEGAL_WINDOWS_CHARS = set('<>:"/\\|?*')
RESERVED_WINDOWS_NAMES = (
    {"CON", "PRN", "AUX", "NUL"}
    | {f"COM{i}" for i in range(1, 10)}
    | {f"LPT{i}" for i in range(1, 10)}
)


# ---------- 資料結構 ----------
@dataclass
class SheetData:
    """單一檔案單一分頁的解析結果。"""

    file_name: str
    sheet_name: str
    headers: list[str] = field(default_factory=list)
    rows: list[list] = field(default_factory=list)
    row_numbers: list[int] = field(default_factory=list)  # 每列在「原始 Excel」的列號
    numeric_sums: dict[str, float] = field(default_factory=dict)
    numeric_counts: dict[str, int] = field(default_factory=dict)
    alerts: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    error: str | None = None


@dataclass
class FileResult:
    """單一檔案的解析結果（多分頁）。"""

    file_name: str
    sheets: list[SheetData] = field(default_factory=list)
    error: str | None = None


# ---------- Windows 路徑預檢（今天學的：pathlib / UNC / MAX_PATH） ----------
def looks_like_windows_path(raw: str) -> bool:
    """判斷字串是否為 Windows 形式的路徑（有磁碟機代號或 UNC 前綴）。

    在 Linux 上 Path('C:/x') 只是普通目錄名，所以判斷 Windows 語意必須用 PureWindowsPath。
    """
    if not raw:
        return False
    wp = PureWindowsPath(raw)
    return bool(wp.drive) or raw.startswith("\\\\") or raw.startswith("//")


def validate_windows_component(name: str) -> list[str]:
    """檢查單一檔名/目錄名是否符合 Windows 命名規則。回傳問題清單（空 = 合法）。"""
    problems: list[str] = []
    bad = sorted({c for c in ILLEGAL_WINDOWS_CHARS if c in name})
    if bad:
        problems.append(
            f"檔名「{name}」含 Windows 禁用字元 {' '.join(bad)}"
            "（< > : \" / \\ | ? * 都不能用，請改用全形或底線）"
        )
    stem = name.split(".")[0].upper()
    if stem in RESERVED_WINDOWS_NAMES:
        problems.append(
            f"檔名「{name}」是 Windows 保留裝置名（{stem}），加副檔名也不行，請改名"
        )
    return problems


def validate_windows_path(raw: str) -> list[str]:
    """對 Windows 形式的路徑做完整預檢：各層名稱 + 總長度 260。"""
    problems: list[str] = []
    wp = PureWindowsPath(raw)
    anchor = wp.anchor  # 'C:\\' 或 '\\\\server\\share\\' —— 這段本身是合法的，不該被當檔名檢查
    for part in wp.parts:
        if part == anchor or part in (wp.drive, "\\", "/"):
            continue
        if part.startswith("\\\\") or part.startswith("//"):
            continue
        problems.extend(validate_windows_component(part))
    if len(str(wp)) > MAX_WINDOWS_PATH:
        problems.append(
            f"路徑長度 {len(str(wp))} 超過 Windows MAX_PATH {MAX_WINDOWS_PATH}；"
            "解法：改用較短路徑、或開啟 LongPathsEnabled + 程式 manifest 標 longPathAware"
        )
    return problems


def validate_target_path(raw: str) -> list[str]:
    """目標輸出路徑預檢：Windows 形式走 Windows 規則，其他走通用規則。"""
    if looks_like_windows_path(raw):
        problems: list[str] = []
        # 跨 OS 守衛：Windows 形式的路徑不能在 Linux/macOS 上執行——
        # 會產生字面檔名「C:\out\merged.xlsx」（一個檔名裡含反斜線的怪檔），
        # 真正的輸出會落在 cwd 而不是預期的磁碟機位置。
        if os.name != "nt":
            problems.append(
                f"輸出路徑「{raw}」是 Windows 形式，但目前作業系統是 {os.name}（非 Windows）。"
                "在非 Windows 上這會被當成含反斜線的單一檔名、寫進目前目錄，"
                "不是預期的磁碟機位置。請改用本機絕對路徑，"
                "或在 Windows 上執行本程式。"
            )
        problems.extend(validate_windows_path(raw))
        return problems
    problems = []
    for part in Path(raw).parts:
        if part in ("/", "\\"):
            continue
        problems.extend(validate_windows_component(part))
    return problems


# ---------- 原子寫入（今天學的：暫存檔 + fsync + os.replace） ----------
def atomic_write_bytes(target: Path, payload: bytes) -> None:
    """把 payload 原子寫入 target：同目錄暫存檔 → flush+fsync → os.replace。

    任一時刻中斷：舊檔完整、新檔不存在，不會出現半套輸出。
    """
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(dir=str(target.parent), prefix=f".{target.name}.", suffix=".tmp")
    tmp_path = Path(tmp_name)
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(payload)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(tmp_path, target)
    except BaseException:
        tmp_path.unlink(missing_ok=True)
        raise


def atomic_save_workbook(wb, target: Path, backup: bool = True) -> Path | None:
    """把 openpyxl Workbook 原子寫出；回傳備份路徑（若有備份）。"""
    target.parent.mkdir(parents=True, exist_ok=True)
    backup_path = None
    if backup and target.exists():
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        # 備份副檔名刻意「不」保留 .xlsx：否則它會被 .xlsx 掃描器選中，
        # 造成重跑時把備份檔當成輸入合併 → 資料重複、金額 double（見 scan_inputs 的排除規則）。
        backup_path = target.with_name(f"{target.stem}.bak-{stamp}")
        shutil.copy2(target, backup_path)

    fd, tmp_name = tempfile.mkstemp(dir=str(target.parent), prefix=f".{target.stem}.", suffix=".xlsx")
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(tmp_path)
        # 必須用 "rb+"（可寫 handle）做 fsync：
        # Windows 的 os.fsync() 底層是 MSVCRT _commit() → FlushFileBuffers()，
        # 而 FlushFileBuffers 要求 handle 具備 GENERIC_WRITE 權限；
        # 用 "rb" 開的 handle 只有 GENERIC_READ → EBADF (OSError [Errno 9] Bad file descriptor)。
        # POSIX 允許對 O_RDONLY fd 做 fsync，所以 Linux CI 抓不到這個 bug（參考 maigret PR #3143）。
        with open(tmp_path, "rb+") as fh:
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(tmp_path, target)
    except BaseException:
        tmp_path.unlink(missing_ok=True)
        raise
    return backup_path


# ---------- 讀取 ----------
def _cell_value(value):
    """把 openpyxl 讀到的值轉成可安全寫回的型別。"""
    from datetime import date, datetime as dt, time as tm

    if value is None:
        return None
    if isinstance(value, (dt, date, tm)):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    return str(value)


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def read_file(path: Path) -> FileResult:
    """讀單一 Excel：逐分頁取表頭、資料列、數值合計、警示。"""
    result = FileResult(file_name=path.name)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 對外 IO 一律要可行動錯誤
        result.error = (
            f"無法開啟 {path.name}：{type(exc).__name__}: {exc}。"
            "常見原因：檔案被 Excel 鎖住（關掉 Excel 再試）、檔案損毀、"
            "或密碼保護（先另存為無密碼版本）。"
        )
        return result

    try:
        for sheet_name in wb.sheetnames:
            sheet = SheetData(file_name=path.name, sheet_name=sheet_name)
            try:
                rows_iter = list(wb[sheet_name].iter_rows(values_only=True))
            except Exception as exc:  # noqa: BLE001
                sheet.error = f"讀取分頁失敗：{type(exc).__name__}: {exc}"
                result.sheets.append(sheet)
                continue

            # 保留「原始 Excel 列號」：過濾空列前先記下來，
            # 否則 _來源列號 會跟原始檔對不上（空列越多偏移越大，來源追溯就不可信）。
            indexed_rows = [
                (orig_idx, r) for orig_idx, r in enumerate(rows_iter, start=1)
            ]
            raw_rows = [(i, r) for i, r in indexed_rows if any(v is not None for v in r)]
            if not raw_rows:
                sheet.error = "分頁是空的"
                result.sheets.append(sheet)
                continue

            header_orig_idx, header_row = raw_rows[0]

            # 有效欄位範圍 = 表頭列最後一個非 None 的位置。
            # openpyxl 的 iter_rows 會把資料列補齊到該分頁最大欄數，表頭尾端因此出現 None。
            # 若把這些 None 命名成「欄4」「欄5」，使用者只會困惑，不知道原始資料其實有那兩欄。
            last_named = -1
            for idx, h in enumerate(header_row):
                if h is not None and str(h).strip():
                    last_named = idx
            if last_named < 0:
                sheet.error = "表頭列沒有任何欄位名稱"
                result.sheets.append(sheet)
                continue

            headers: list[str] = []
            used: dict[str, int] = {}
            for idx in range(last_named + 1):
                h = header_row[idx]
                name = "" if h is None else str(h).strip()
                if not name:
                    name = f"欄{idx + 1}"
                # 重複欄名自動去重：原本 dict(zip(headers, values)) 會讓後者覆蓋前者 → 資料無聲遺失
                if name in used:
                    used[name] += 1
                    new_name = f"{name}_{used[name]}"
                    sheet.warnings.append(
                        f"表頭有重複欄名「{name}」（第 {idx + 1} 欄），已自動改名為「{new_name}」"
                    )
                    name = new_name
                else:
                    used[name] = 1
                headers.append(name)
            sheet.headers = headers

            for r_idx, raw in raw_rows[1:]:
                values = [_cell_value(v) for v in raw]
                if len(values) > len(headers):
                    sheet.warnings.append(
                        f"資料列比表頭多 {len(values) - len(headers)} 欄，已截斷"
                    )
                values += [None] * (len(headers) - len(values))
                values = values[: len(headers)]
                sheet.rows.append(values)
                sheet.row_numbers.append(r_idx)

                for c_idx, value in enumerate(values):
                    header = headers[c_idx]
                    if _is_number(value):
                        sheet.numeric_sums[header] = sheet.numeric_sums.get(header, 0.0) + float(value)
                        sheet.numeric_counts[header] = sheet.numeric_counts.get(header, 0) + 1

                    reasons: list[str] = []
                    header_lower = header.lower()
                    if any(kw.lower() in header_lower for kw in REFUND_KEYWORDS):
                        reasons.append("退款關鍵欄")
                    if _is_number(value) and float(value) < 0:
                        reasons.append("負值")
                    if reasons:
                        sheet.alerts.append(
                            {
                                "file": path.name,
                                "sheet": sheet_name,
                                "row": r_idx,
                                "header": header,
                                "value": value,
                                "reason": "、".join(reasons),
                            }
                        )
            result.sheets.append(sheet)
    finally:
        wb.close()
    return result


# ---------- 合併 ----------
def scan_inputs(folder: Path, excluded: set[Path]) -> tuple[list[Path], list[tuple[str, str]]]:
    """列出要合併的 .xlsx。回傳 (要合併的檔案, 被略過的檔案[(檔名, 原因)])。

    略過的檔案要回報——靜默略過會讓使用者困惑「我的檔怎麼沒被合併」。

    `excluded` 是「自己產生的輸出檔」集合（輸出檔本身 + 它的時間戳備份 + CSV 摘要），
    必須排除，否則重跑時會把上一次的輸出/備份當成新的輸入合併 → 資料重複、金額 double。
    """
    files: list[Path] = []
    skipped: list[tuple[str, str]] = []
    excluded_resolved = set()
    for e in excluded:
        try:
            excluded_resolved.add(e.resolve())
        except OSError:
            excluded_resolved.add(e)
    for path in sorted(folder.iterdir()):
        if not path.is_file():
            continue
        if path.name.startswith(SKIP_PREFIXES):
            continue
        suffix = path.suffix.lower()
        if suffix not in INPUT_SUFFIXES:
            if suffix in OTHER_EXCEL_SUFFIXES:
                skipped.append(
                    (path.name, f"副檔名 {suffix} 不支援（{suffix} 舊格式請在 Excel 另存為 .xlsx）")
                )
            continue
        try:
            resolved = path.resolve()
        except OSError:
            resolved = path
        if resolved in excluded_resolved:
            skipped.append((path.name, "是輸出檔本身，已排除（避免自我合併）"))
            continue
        files.append(path)
    return files, skipped


def excluded_artifacts(output_path: Path, csv_path: Path | None = None) -> set[Path]:
    """列出「這個工具自己產生的檔案」：輸出檔本體 + 時間戳備份 + 摘要 CSV。

    備份檔名格式為 <stem>.bak-<YYYYMMDD-HHMMSS>（不保留 .xlsx 副檔名，
    所以理論上不會被 .xlsx 掃描器選中）。但仍在此明確排除，
    以防使用者手上已有舊版（v2 舊碼）產生的 <stem>.bak-*.xlsx 備份檔。
    """
    out: set[Path] = {output_path}
    stem = output_path.stem
    for pat in (f"{stem}.bak-*", f".{stem}.bak-*", f"{stem}.bak"):
        out.update(output_path.parent.glob(pat))
    if csv_path is not None:
        out.add(csv_path)
        out.update(csv_path.parent.glob(f"{csv_path.stem}.bak-*"))
    return out


def union_headers(sheets: list[SheetData]) -> list[str]:
    """跨檔欄位聯集：依首次出現順序。"""
    out: list[str] = []
    seen: set[str] = set()
    for sheet in sheets:
        for h in sheet.headers:
            if h not in seen:
                seen.add(h)
                out.append(h)
    return out


def merge_all(
    folder: Path, output_path: Path, csv_path: Path | None = None
) -> tuple[list[FileResult], list[SheetData], list[str], list[tuple[str, str]]]:
    """掃描並解析資料夾內所有 Excel。

    回傳 (各檔結果, 可用分頁, 錯誤清單, 略過檔案[(檔名, 原因)])。
    """
    results: list[FileResult] = []
    usable: list[SheetData] = []
    errors: list[str] = []
    # 排除「這個工具自己產生的檔案」：輸出檔本體 + 時間戳備份 + 摘要 CSV。
    # 不排除備份 → 重跑時備份檔被當輸入合併 → 資料重複、金額 double（實測 8 列 → 36 列）。
    excluded = excluded_artifacts(output_path, csv_path)
    files, skipped = scan_inputs(folder, excluded)

    for path in files:
        result = read_file(path)
        results.append(result)
        if result.error:
            errors.append(result.error)
            print(f"  ✗ 略過 {path.name}：{result.error}", file=sys.stderr)
            continue
        for sheet in result.sheets:
            if sheet.error:
                errors.append(f"{path.name} / {sheet.sheet_name}：{sheet.error}")
                print(f"  ✗ {path.name} / {sheet.sheet_name}：{sheet.error}", file=sys.stderr)
                continue
            usable.append(sheet)
    return results, usable, errors, skipped


# ---------- 寫出 ----------
def _style_header(ws, headers: list[str], fill: str = "D9E1F2") -> None:
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)


def _autosize(ws, widths: list[int]) -> None:
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w


def build_workbook(
    results: list[FileResult],
    usable: list[SheetData],
    union: list[str],
    input_folder: Path,
    output_path: Path,
    csv_path: Path | None,
    csv_encoding: str,
    skipped: list[tuple[str, str]] | None = None,
) -> Workbook:
    wb = Workbook()

    # Sheet 1：合併資料（欄位聯集 + 來源追溯欄）
    ws_data = wb.active
    ws_data.title = "合併資料"
    headers = union + ["_來源檔", "_來源分頁", "_來源列號"]
    _style_header(ws_data, headers)
    row_idx = 2
    for sheet in usable:
        for i, values in enumerate(sheet.rows):
            mapping = dict(zip(sheet.headers, values))
            cells = [mapping.get(h) for h in union]
            # 用「原始 Excel 列號」而不是過濾後的索引，否則來源追溯對不上原始檔
            src_row = sheet.row_numbers[i] if i < len(sheet.row_numbers) else i + 2
            cells += [sheet.file_name, sheet.sheet_name, src_row]
            for c_idx, value in enumerate(cells, start=1):
                ws_data.cell(row=row_idx, column=c_idx, value=value)
            row_idx += 1
    _autosize(ws_data, [18] * len(union) + [24, 16, 10])

    # Sheet 2：檔案摘要
    ws_meta = wb.create_sheet("檔案摘要")
    meta_headers = ["檔名", "分頁", "資料列數", "欄位數", "狀態", "錯誤訊息", "缺漏欄位", "獨有欄位", "警告"]
    _style_header(ws_meta, meta_headers)
    row_idx = 2
    for result in results:
        if result.error:
            ws_meta.cell(row=row_idx, column=1, value=result.file_name)
            ws_meta.cell(row=row_idx, column=5, value="失敗")
            ws_meta.cell(row=row_idx, column=6, value=result.error)
            row_idx += 1
            continue
        for sheet in result.sheets:
            others: set[str] = set()
            for other in result.sheets:
                if other is not sheet:
                    others |= set(other.headers)
            for other_result in results:
                if other_result is result:
                    continue
                for other_sheet in other_result.sheets:
                    others |= set(other_sheet.headers)
            missing = [h for h in union if h not in sheet.headers]
            extra = [h for h in sheet.headers if h not in others]
            ws_meta.cell(row=row_idx, column=1, value=result.file_name)
            ws_meta.cell(row=row_idx, column=2, value=sheet.sheet_name)
            ws_meta.cell(row=row_idx, column=3, value=len(sheet.rows))
            ws_meta.cell(row=row_idx, column=4, value=len(sheet.headers))
            status = "失敗" if sheet.error else ("只有表頭" if not sheet.rows else "成功")
            ws_meta.cell(row=row_idx, column=5, value=status)
            ws_meta.cell(row=row_idx, column=6, value=sheet.error or "")
            ws_meta.cell(row=row_idx, column=7, value="、".join(missing))
            ws_meta.cell(row=row_idx, column=8, value="、".join(extra))
            ws_meta.cell(row=row_idx, column=9, value="；".join(sheet.warnings))
            row_idx += 1
    _autosize(ws_meta, [26, 18, 10, 10, 10, 50, 30, 30, 40])

    # Sheet 3：數值合計
    ws_sum = wb.create_sheet("數值合計")
    sum_headers = ["來源檔", "分頁", "欄位", "合計", "非空數值格數"]
    _style_header(ws_sum, sum_headers)
    row_idx = 2
    for sheet in usable:
        for header, total in sheet.numeric_sums.items():
            ws_sum.cell(row=row_idx, column=1, value=sheet.file_name)
            ws_sum.cell(row=row_idx, column=2, value=sheet.sheet_name)
            ws_sum.cell(row=row_idx, column=3, value=header)
            ws_sum.cell(row=row_idx, column=4, value=total)
            ws_sum.cell(row=row_idx, column=5, value=sheet.numeric_counts.get(header, 0))
            row_idx += 1
    _autosize(ws_sum, [26, 18, 22, 14, 14])

    # Sheet 4：退款/負值警示
    ws_alert = wb.create_sheet("退款警示")
    alert_headers = ["來源檔", "分頁", "列號", "欄位", "值", "原因"]
    _style_header(ws_alert, alert_headers, fill="FCE4D6")
    row_idx = 2
    for sheet in usable:
        for alert in sheet.alerts:
            ws_alert.cell(row=row_idx, column=1, value=alert["file"])
            ws_alert.cell(row=row_idx, column=2, value=alert["sheet"])
            ws_alert.cell(row=row_idx, column=3, value=alert["row"])
            ws_alert.cell(row=row_idx, column=4, value=alert["header"])
            ws_alert.cell(row=row_idx, column=5, value=alert["value"])
            ws_alert.cell(row=row_idx, column=6, value=alert["reason"])
            row_idx += 1
    if row_idx == 2:
        ws_alert.cell(row=2, column=1, value="（無退款/負值紀錄）")
    _autosize(ws_alert, [26, 18, 8, 22, 14, 20])

    # Sheet 5：本次執行
    ws_run = wb.create_sheet("本次執行")
    run_rows = [
        ("執行時間", datetime.now().isoformat(timespec="seconds")),
        ("輸入資料夾", str(input_folder)),
        ("輸出檔", str(output_path)),
        ("摘要 CSV", str(csv_path) if csv_path else "（未輸出）"),
        ("CSV 編碼", csv_encoding if csv_path else "-"),
        ("解析檔數", len(results)),
        ("可用分頁數", len(usable)),
        ("合併後總列數", sum(len(s.rows) for s in usable)),
        ("聯集欄位數", len(union)),
        ("警示筆數", sum(len(s.alerts) for s in usable)),
        ("略過檔案數", len(skipped or [])),
    ]
    for r_idx, (k, v) in enumerate(run_rows, start=1):
        ws_run.cell(row=r_idx, column=1, value=k).font = Font(bold=True)
        ws_run.cell(row=r_idx, column=2, value=v)
    _autosize(ws_run, [18, 60])

    # Sheet 6：略過的檔案（靜默略過會讓使用者困惑「我的檔怎麼沒被合併」）
    ws_skip = wb.create_sheet("略過檔案")
    skip_headers = ["檔名", "略過原因"]
    _style_header(ws_skip, skip_headers, fill="FFF2CC")
    row_idx = 2
    for name, reason in (skipped or []):
        ws_skip.cell(row=row_idx, column=1, value=name)
        ws_skip.cell(row=row_idx, column=2, value=reason)
        row_idx += 1
    if row_idx == 2:
        ws_skip.cell(row=2, column=1, value="（無略過檔案）")
    _autosize(ws_skip, [30, 70])

    return wb


def summary_rows(results: list[FileResult], usable: list[SheetData], union: list[str]) -> list[list]:
    """組出 CSV 摘要列（與工作表的「檔案摘要」一致 —— 欄位必須一一對應）。"""
    rows: list[list] = [[
        "檔名", "分頁", "資料列數", "欄位數", "狀態", "錯誤訊息", "缺漏欄位", "獨有欄位", "警告"
    ]]
    for result in results:
        if result.error:
            rows.append([result.file_name, "", "", "", "失敗", result.error, "", "", ""])
            continue
        for sheet in result.sheets:
            others: set[str] = set()
            for other_sheet in result.sheets:
                if other_sheet is not sheet:
                    others |= set(other_sheet.headers)
            for other_result in results:
                if other_result is result:
                    continue
                for other_sheet in other_result.sheets:
                    others |= set(other_sheet.headers)
            missing = "、".join(h for h in union if h not in sheet.headers)
            extra = "、".join(h for h in sheet.headers if h not in others)
            status = "失敗" if sheet.error else ("只有表頭" if not sheet.rows else "成功")
            rows.append(
                [
                    result.file_name,
                    sheet.sheet_name,
                    len(sheet.rows),
                    len(sheet.headers),
                    status,
                    sheet.error or "",
                    missing,
                    extra,
                    "；".join(sheet.warnings),
                ]
            )
    return rows


def atomic_write_csv(target: Path, rows: list[list], encoding: str) -> None:
    """原子寫出 CSV（預設 utf-8-sig，Excel 開不亂碼）。"""
    import io

    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    for row in rows:
        writer.writerow(row)
    payload = buffer.getvalue().encode(encoding, errors="replace")
    atomic_write_bytes(target, payload)


# ---------- CLI ----------
def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="合併資料夾內多個 Excel 並產生摘要（v2：原子寫入 + CSV BOM + Windows 路徑預檢）",
    )
    parser.add_argument("input_folder", type=Path, help="要合併的資料夾")
    parser.add_argument(
        "-o", "--output", type=Path, default=Path("merged.xlsx"), help="輸出 Excel（預設 merged.xlsx）"
    )
    parser.add_argument("--summary-csv", type=Path, default=None, help="另存摘要 CSV 的路徑")
    parser.add_argument(
        "--csv-encoding", choices=CSV_ENCODINGS, default="utf-8-sig",
        help="摘要 CSV 編碼（預設 utf-8-sig：Excel 直接開不亂碼；cp950 給舊系統）",
    )
    parser.add_argument(
        "--apply", action="store_true",
        help="真的寫出檔案（預設 dry-run：只預覽不寫）",
    )
    parser.add_argument("--no-backup", action="store_true", help="覆寫前不建立時間戳備份")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)

    if not args.input_folder.is_dir():
        sys.stderr.write(
            f"錯誤：{args.input_folder} 不是資料夾。請確認路徑是否存在，"
            "Windows 下可用引號包住含空格的路徑。\n"
        )
        return 2

    path_problems = validate_target_path(str(args.output))
    if args.summary_csv:
        path_problems += validate_target_path(str(args.summary_csv))
    if path_problems:
        sys.stderr.write("輸出路徑預檢失敗（先修正再執行，避免在 Windows 上產生無法開啟的檔名）：\n")
        for p in path_problems:
            sys.stderr.write(f"  - {p}\n")
        return 2

    print(f"[merge_xlsx] 輸入資料夾：{args.input_folder}")
    print(f"[merge_xlsx] 輸出 Excel：{args.output}")
    if args.summary_csv:
        print(f"[merge_xlsx] 摘要 CSV：{args.summary_csv}（編碼 {args.csv_encoding}）")
    print(f"[merge_xlsx] 模式：{'APPLY（會寫檔）' if args.apply else 'dry-run（不寫檔）'}")

    results, usable, errors, skipped = merge_all(args.input_folder, args.output, args.summary_csv)

    if not results:
        sys.stderr.write(
            f"錯誤：{args.input_folder} 內找不到 .xlsx 檔案。"
            "請確認資料夾正確、檔案副檔名是 .xlsx（.xls 舊格式請先在 Excel 另存為 .xlsx）。\n"
        )
        return 2
    if not usable:
        sys.stderr.write("錯誤：所有檔案都解析失敗，沒有可合併的資料。\n")
        for e in errors:
            sys.stderr.write(f"  - {e}\n")
        return 4

    union = union_headers(usable)
    print(f"  解析成功：{sum(1 for r in results if not r.error)} 個檔案 / {len(usable)} 個分頁")
    print(f"  聯集欄位（{len(union)}）：{'、'.join(union)}")
    print(f"  合併後總列數：{sum(len(s.rows) for s in usable)}")
    print(f"  退款/負值警示：{sum(len(s.alerts) for s in usable)} 筆")
    if skipped:
        print(f"  略過 {len(skipped)} 個檔案（原因見輸出檔的「略過檔案」分頁）：")
        for name, reason in skipped:
            print(f"    - {name}：{reason}")

    if not args.apply:
        print("\n[dry-run] 沒有寫任何檔案。要真的輸出請加 --apply。")
        return 0

    wb = build_workbook(
        results, usable, union, args.input_folder, args.output, args.summary_csv,
        args.csv_encoding, skipped,
    )
    # 寫檔是唯一會壞掉資料的步驟 → 必須給可行動訊息，不能裸 traceback。
    # exit code 3 = 輸出失敗（2 = 用法/路徑錯誤、4 = 資料全壞）。
    try:
        backup = atomic_save_workbook(wb, args.output, backup=not args.no_backup)
    except PermissionError:
        sys.stderr.write(
            f"錯誤：沒有權限寫入「{args.output}」。\n"
            "請改用有寫入權限的資料夾，或用 -o 指定其他輸出路徑。\n"
        )
        return 3
    except OSError as exc:
        sys.stderr.write(
            f"錯誤：寫出「{args.output}」失敗（{type(exc).__name__}: {exc}）。\n"
            "請確認磁碟空間、檔案是否被 Excel 開著鎖住、以及輸出資料夾是否存在。\n"
        )
        return 3
    print(f"  ✓ 寫出 {args.output}")
    if backup:
        print(f"  ↻ 舊檔已備份到 {backup}")

    if args.summary_csv:
        try:
            atomic_write_csv(args.summary_csv, summary_rows(results, usable, union), args.csv_encoding)
        except PermissionError:
            sys.stderr.write(
                f"錯誤：沒有權限寫入摘要 CSV「{args.summary_csv}」。\n"
                "Excel 已輸出完成，但 CSV 摘要未產生（可用 --summary-csv 換個路徑重跑）。\n"
            )
            return 3
        except OSError as exc:
            sys.stderr.write(
                f"錯誤：寫出摘要 CSV「{args.summary_csv}」失敗（{type(exc).__name__}: {exc}）。\n"
                "Excel 已輸出完成，但 CSV 摘要未產生。\n"
            )
            return 3
        print(f"  ✓ 寫出 {args.summary_csv}（{args.csv_encoding}）")

    return 0


if __name__ == "__main__":
    sys.exit(main())
