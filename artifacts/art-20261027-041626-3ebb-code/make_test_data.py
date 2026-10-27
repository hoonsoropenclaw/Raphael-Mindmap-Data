#!/usr/bin/env python3
"""產生 v4 的測試資料：涵蓋六項改進的邊界案例。

設計原則（SOP §2）：至少 3 組輸入（正常／空缺資料／異常格式），
而且測試資料本身要先驗證有效 —— 每個案例對應一個已知預期結果。
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook


def _save(wb: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_normal(path: Path) -> None:
    """正常檔：表頭在第 1 列，數值是原生型別。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "經費"
    ws.append(["經費項目", "核定金額", "執行率"])
    ws.append(["業務費", 500000, 0.95])
    ws.append(["設備費", 300000, 0.88])
    ws.append(["合計", 800000, None])
    _save(wb, path)


def make_preface(path: Path) -> None:
    """表頭不在第 1 列：前有說明列（型別互斥判準的關鍵案例）。

    v3 在這個檔會把「本表為 112 年度結算表」當成表頭 → 整批欄位名錯掉。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "112年度"
    ws.append(["本表為 112 年度教育處結算表，單位：元"])
    ws.append([])  # 空列（測試「空列不計入」）
    ws.append(["經費項目", "核定金額", "執行率"])
    ws.append(["業務費", 400000, 0.9])
    ws.append(["旅運費", 100000, 0.8])
    _save(wb, path)


def make_text_numbers(path: Path) -> None:
    """文字型數字：金額是文字、千分位、幣別、百分比、全形、括號負數。

    v3 的 _is_number 只認 int/float → 這些值 100% 靜默漏算。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "經費"
    ws.append(["經費項目", "核定金額", "執行率"])
    ws.append(["業務費", "1,234", "45%"])
    ws.append(["設備費", "NT$2,000", "0.5"])
    ws.append(["差旅費", "１，２３４．５", "30%"])
    ws.append(["退款", "(500)", "0"])
    _save(wb, path)


def make_unparseable(path: Path) -> None:
    """無法解析：金額欄位混了非數字值（必須進「無法解析」清單，不算 0、不丟棄）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "經費"
    ws.append(["經費項目", "核定金額"])
    ws.append(["業務費", "約一百萬"])
    ws.append(["設備費", 200000])
    ws.append(["其他", "待補"])
    _save(wb, path)


def make_merged_header(path: Path) -> None:
    """合併儲存格：A1:B1 合併 → 欄名被吞掉一個（要自己補回來）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "經費"
    ws["A1"] = "經費項目"
    ws["B1"] = "核定金額"
    ws.merge_cells("A1:B1")
    ws.append(["業務費", 100000])
    ws.append(["設備費", 200000])
    _save(wb, path)


def make_rate_semantics(path: Path) -> None:
    """百分比語意要靠欄位語意決定：件數欄出現 '45%' 必須報無法解析（不猜）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "經費"
    ws.append(["經費項目", "核定金額", "件數"])
    ws.append(["業務費", 100000, "45%"])
    ws.append(["設備費", 200000, 3])
    _save(wb, path)


def make_empty(path: Path) -> None:
    """空檔：完全沒有資料列（只有表頭）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "空表"
    ws.append(["經費項目", "核定金額"])
    _save(wb, path)


def make_multi_sheet(path: Path) -> None:
    """多分頁，且兩個分頁的表頭列位置不同（每個 sheet 獨立判定）。"""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "第一季"
    ws1.append(["經費項目", "核定金額"])
    ws1.append(["業務費", 100000])

    ws2 = wb.create_sheet("第二季")
    ws2.append(["說明：第二季含追加預算"])
    ws2.append(["經費項目", "核定金額"])
    ws2.append(["設備費", 250000])
    _save(wb, path)


def make_old_format(path: Path) -> None:
    """舊格式副檔名（.xls）：要略過並回報原因。"""
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")  # 只是假檔案


def make_lockfile(path: Path) -> None:
    """Excel 鎖檔（~$xxx）：要被排除。"""
    path.write_bytes(b"lock")


def main(argv: list[str]) -> int:
    out = Path(argv[1]) if len(argv) > 1 else Path("testdata")
    out.mkdir(parents=True, exist_ok=True)
    make_normal(out / "a_正常.xlsx")
    make_preface(out / "b_有說明列.xlsx")
    make_text_numbers(out / "c_文字型數字.xlsx")
    make_unparseable(out / "d_無法解析.xlsx")
    make_merged_header(out / "e_合併儲存格.xlsx")
    make_rate_semantics(out / "f_百分比語意.xlsx")
    make_empty(out / "g_只有表頭.xlsx")
    make_multi_sheet(out / "h_多分頁.xlsx")
    make_old_format(out / "i_舊格式.xls")
    make_lockfile(out / "~$j_鎖檔.xlsx")
    print(f"測試資料已產生於 {out.resolve()}")
    for p in sorted(out.iterdir()):
        print(f"  - {p.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
