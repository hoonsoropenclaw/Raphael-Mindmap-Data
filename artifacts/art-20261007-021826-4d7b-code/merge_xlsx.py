#!/usr/bin/env python3
"""
merge_xlsx.py — 把資料夾內多個 Excel 合併並產生摘要。

功能：
  1. 掃描 input_dir 內所有 .xlsx
  2. 用 pandas 讀入，標記來源檔；缺值視為 0
  3. 合併成單一 DataFrame，寫到 output.xlsx 的「原始合併」工作表
  4. 摘要工作表：總筆數、總金額、各檔筆數與金額、Top 5 項目
  5. 加一張金額長條圖（每檔總金額）
  6. 失敗時：印出可行動錯誤訊息、退出碼 1、保留之前的輸出不被覆蓋（自動加 .bak）

執行：
  python merge_xlsx.py --input <dir> --output <path.xlsx>

設計原則（SOP §1）：
  - 先寫最小可跑，再加錯誤處理與圖表
  - 破壞性動作（覆寫 output）預設先備份原檔
  - 不刪檔、不寄信、不寫死金鑰
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference


REQUIRED_COLS = {"項目", "日期", "金額", "備註"}


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="合併資料夾內多個 Excel 並產生摘要")
    p.add_argument("--input", "-i", required=True, type=Path, help="輸入資料夾")
    p.add_argument("--output", "-o", required=True, type=Path, help="輸出 xlsx 路徑")
    p.add_argument("--no-backup", action="store_true", help="覆寫輸出時不備份")
    return p.parse_args()


def find_inputs(input_dir: Path) -> list[Path]:
    """掃所有 .xlsx，但跳過以 ~$ 開頭（Excel 暫存檔）"""
    if not input_dir.exists():
        raise FileNotFoundError(f"找不到輸入資料夾: {input_dir}")
    if not input_dir.is_dir():
        raise NotADirectoryError(f"不是資料夾: {input_dir}")
    files = sorted(p for p in input_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        raise ValueError(f"資料夾內沒有 .xlsx: {input_dir}")
    return files


def read_one(path: Path) -> pd.DataFrame:
    """讀一個檔；缺欄位或解析錯誤給可行動提示"""
    try:
        df = pd.read_excel(path, dtype={"金額": "float64"})
    except Exception as e:
        raise RuntimeError(f"讀取 {path.name} 失敗（檔案可能損壞或非 xlsx 格式）: {e}") from e

    missing = REQUIRED_COLS - set(df.columns)
    if missing:
        raise ValueError(f"{path.name} 缺少必要欄位: {sorted(missing)}（需要 {sorted(REQUIRED_COLS)}）")

    df = df.copy()
    df["來源檔"] = path.name
    df["金額"] = df["金額"].fillna(0)  # 缺值計 0，不算入錯誤
    return df


def merge_all(files: list[Path]) -> pd.DataFrame:
    dfs = [read_one(p) for p in files]
    combined = pd.concat(dfs, ignore_index=True)
    return combined


def summarize(df: pd.DataFrame) -> dict:
    by_src = df.groupby("來源檔").agg(筆數=("金額", "size"), 總金額=("金額", "sum"))
    by_src = by_src.sort_values("總金額", ascending=False)
    top5 = df.nlargest(5, "金額")[["項目", "來源檔", "日期", "金額"]]
    return {
        "總筆數": int(len(df)),
        "總金額": float(df["金額"].sum()),
        "來源檔數": int(df["來源檔"].nunique()),
        "by_src": by_src,
        "top5": top5.reset_index(drop=True),
    }


def write_output(df: pd.DataFrame, summary: dict, output: Path) -> None:
    """寫兩個工作表：原始合併 + 摘要；摘要頁加長條圖。"""
    output.parent.mkdir(parents=True, exist_ok=True)

    # 1) 先用 pandas 寫原始合併（最簡單可靠）
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="原始合併", index=False)
        s = summary
        pd.DataFrame(
            [{"指標": "總筆數", "值": s["總筆數"]},
             {"指標": "總金額", "值": s["總金額"]},
             {"指標": "來源檔數", "值": s["來源檔數"]}]
        ).to_excel(writer, sheet_name="摘要", startrow=0, index=False)
        s["by_src"].to_excel(writer, sheet_name="摘要", startrow=5)
        s["top5"].to_excel(writer, sheet_name="摘要", startrow=len(s["by_src"]) + 8)

    # 2) openpyxl 加圖（pd.ExcelWriter 已開檔，load 再寫）
    wb = load_workbook(output)
    ws = wb["摘要"]
    chart = BarChart()
    chart.title = "各來源檔總金額"
    chart.y_axis.title = "金額"
    chart.x_axis.title = "來源檔"
    # by_src 從第 6 列（index=5）開始，A 欄是來源檔名、B 欄是筆數、C 欄是總金額
    by_src_start = 6
    by_src_end = by_src_start + len(s["by_src"])
    data = Reference(ws, min_col=3, min_row=by_src_start, max_row=by_src_end)  # 總金額欄
    cats = Reference(ws, min_col=1, min_row=by_src_start + 1, max_row=by_src_end)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 10
    chart.width = 18
    ws.add_chart(chart, f"E{by_src_start}")
    wb.save(output)


def backup_if_exists(path: Path, enabled: bool) -> None:
    if path.exists() and enabled:
        bak = path.with_suffix(path.suffix + ".bak")
        shutil.copy2(path, bak)
        print(f"已備份原檔 → {bak.name}")


def main() -> int:
    args = parse_args()
    try:
        files = find_inputs(args.input)
        print(f"找到 {len(files)} 個檔案:")
        for f in files:
            print(f"  - {f.name}")

        df = merge_all(files)
        summary = summarize(df)
        print("\n=== 摘要 ===")
        print(f"總筆數: {summary['總筆數']}, 總金額: {summary['總金額']:.0f}")
        print(summary["by_src"].to_string())

        backup_if_exists(args.output, not args.no_backup)
        write_output(df, summary, args.output)
        print(f"\n寫入完成: {args.output}")
        return 0
    except FileNotFoundError as e:
        print(f"錯誤：{e}", file=sys.stderr)
        return 1
    except (NotADirectoryError, ValueError) as e:
        print(f"錯誤：{e}", file=sys.stderr)
        return 1
    except RuntimeError as e:
        print(f"錯誤：{e}", file=sys.stderr)
        return 1
    except Exception as e:
        print(f"未預期錯誤：{type(e).__name__}: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())